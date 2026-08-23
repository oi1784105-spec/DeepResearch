"""Private, per-user attachment storage and context extraction for research requests."""

import base64
import io
import json
import mimetypes
import re
from pathlib import Path
from urllib.parse import quote, urljoin
from uuid import uuid4

import httpx
from fastapi import HTTPException, UploadFile, status

from backend.database import db
from backend.schemas.attachment import AttachmentResponse
from backend.security import decrypt_api_key


PROJECT_ROOT = Path(__file__).resolve().parents[3]
UPLOAD_ROOT = PROJECT_ROOT / "data" / "uploads"
MAX_FILES_PER_REQUEST = 5
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS_PER_FILE = 50_000
MAX_CONTEXT_CHARS = 100_000

IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml",
    ".html", ".htm", ".py", ".js", ".ts", ".vue", ".java", ".go",
    ".c", ".cpp", ".h", ".sql", ".log", ".ini", ".toml", ".rst",
}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx"}
SUPPORTED_EXTENSIONS = set(IMAGE_TYPES) | TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS


def _safe_filename(filename: str | None) -> str:
    value = Path(filename or "attachment").name
    value = re.sub(r"[^A-Za-z0-9._()\-\u4e00-\u9fff ]", "_", value).strip(" .")
    return value[:120] or "attachment"


def _trim_text(value: str) -> str:
    value = value.replace("\x00", "").strip()
    if len(value) > MAX_TEXT_CHARS_PER_FILE:
        return value[:MAX_TEXT_CHARS_PER_FILE] + "\n\n[内容过长，已截取前 50,000 个字符]"
    return value


def _read_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return _trim_text(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise HTTPException(status_code=415, detail="无法读取此文本文件的编码，请保存为 UTF-8 后重试")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise HTTPException(status_code=501, detail="服务器尚未安装 PDF 解析组件") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return _trim_text("\n\n".join(page.extract_text() or "" for page in reader.pages))
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"PDF 解析失败：{exc}") from exc


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(status_code=501, detail="服务器尚未安装 Word 解析组件") from exc
    try:
        document = Document(io.BytesIO(data))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                blocks.append("\t".join(cell.text.strip() for cell in row.cells))
        return _trim_text("\n".join(blocks))
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Word 文档解析失败：{exc}") from exc


def _extract_xlsx(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=501, detail="服务器尚未安装 Excel 解析组件") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks: list[str] = []
        for sheet in workbook.worksheets[:5]:
            blocks.append(f"[工作表：{sheet.title}]")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 500:
                    blocks.append("[该工作表已截取前 500 行]")
                    break
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    blocks.append("\t".join(values))
        workbook.close()
        return _trim_text("\n".join(blocks))
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Excel 文件解析失败：{exc}") from exc


def _extract_text(data: bytes, suffix: str) -> str:
    if suffix in TEXT_EXTENSIONS:
        return _read_text(data)
    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".docx":
        return _extract_docx(data)
    if suffix == ".xlsx":
        return _extract_xlsx(data)
    return ""


def _response(row: dict) -> AttachmentResponse:
    return AttachmentResponse(
        id=row["id"],
        filename=row["filename"],
        content_type=row["content_type"],
        size_bytes=int(row["size_bytes"]),
        kind=row["kind"],
        is_image=row["kind"] == "image",
    )


async def save_attachments(files: list[UploadFile], user_id: int) -> list[AttachmentResponse]:
    if not files:
        raise HTTPException(status_code=422, detail="请选择至少一个文件")
    if len(files) > MAX_FILES_PER_REQUEST:
        raise HTTPException(status_code=422, detail=f"一次最多上传 {MAX_FILES_PER_REQUEST} 个文件")

    prepared: list[dict] = []
    for upload in files:
        filename = _safe_filename(upload.filename)
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise HTTPException(status_code=415, detail="仅支持图片、文本、PDF、DOCX 和 XLSX 文件")
        data = await upload.read(MAX_FILE_BYTES + 1)
        if not data:
            raise HTTPException(status_code=422, detail=f"文件为空：{filename}")
        if len(data) > MAX_FILE_BYTES:
            raise HTTPException(status_code=413, detail=f"文件超过 {MAX_FILE_BYTES // 1024 // 1024}MB：{filename}")
        content_type = IMAGE_TYPES.get(suffix) or upload.content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
        kind = "image" if suffix in IMAGE_TYPES else "document"
        prepared.append({
            "id": uuid4().hex,
            "filename": filename,
            "content_type": content_type,
            "size_bytes": len(data),
            "kind": kind,
            "data": data,
            "extracted_text": "" if kind == "image" else _extract_text(data, suffix),
        })

    directory = UPLOAD_ROOT / str(user_id)
    directory.mkdir(parents=True, exist_ok=True)
    stored: list[AttachmentResponse] = []
    with db() as conn:
        for item in prepared:
            path = directory / f"{item['id']}_{item['filename']}"
            path.write_bytes(item.pop("data"))
            item["storage_path"] = str(path.resolve())
            conn.execute(
                "INSERT INTO attachments(id, user_id, filename, content_type, size_bytes, kind, storage_path, extracted_text) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (item["id"], user_id, item["filename"], item["content_type"], item["size_bytes"], item["kind"], item["storage_path"], item["extracted_text"]),
            )
            stored.append(_response(item))
    return stored


def delete_attachment(attachment_id: str, user_id: int) -> None:
    with db() as conn:
        row = conn.execute("SELECT storage_path FROM attachments WHERE id = ? AND user_id = ?", (attachment_id, user_id)).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="附件不存在或无权访问")
        conn.execute("DELETE FROM attachments WHERE id = ? AND user_id = ?", (attachment_id, user_id))
    path = Path(row["storage_path"]).resolve()
    if UPLOAD_ROOT.resolve() in path.parents:
        path.unlink(missing_ok=True)


def delete_attachments(attachment_ids: list[str], user_id: int) -> int:
    """Delete a batch of this user's temporary attachments and their files."""
    unique_ids = list(dict.fromkeys(item.strip() for item in attachment_ids if item and item.strip()))
    if not unique_ids:
        return 0
    placeholders = ",".join("?" for _ in unique_ids)
    with db() as conn:
        rows = conn.execute(
            f"SELECT id, storage_path FROM attachments WHERE user_id = ? AND id IN ({placeholders})",
            (user_id, *unique_ids),
        ).fetchall()
        conn.execute(
            f"DELETE FROM attachments WHERE user_id = ? AND id IN ({placeholders})",
            (user_id, *unique_ids),
        )

    upload_root = UPLOAD_ROOT.resolve()
    for row in rows:
        path = Path(row["storage_path"]).resolve()
        if upload_root in path.parents:
            path.unlink(missing_ok=True)
    return len(rows)


def _endpoint(endpoint: str, suffix: str) -> str:
    base = endpoint.rstrip("/") + "/"
    if endpoint.rstrip("/").endswith(suffix.lstrip("/")):
        return endpoint
    return urljoin(base, suffix.lstrip("/"))


def _message_content(value: object) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return "\n".join(str(item.get("text", "")) for item in value if isinstance(item, dict)).strip()
    return str(value or "").strip()


async def _describe_image(row: dict, provider: dict) -> str:
    path = Path(row["storage_path"])
    if not path.is_file():
        raise HTTPException(status_code=410, detail=f"附件文件已不存在：{row['filename']}")
    data = path.read_bytes()
    encoded = base64.b64encode(data).decode("ascii")
    api_key = decrypt_api_key(provider["api_key_encrypted"])
    endpoint = provider["endpoint"].rstrip("/")
    instruction = "请准确识别这张图片中的文字、表格、图表、关键对象和数据。只输出可供后续研究使用的事实性摘要，不要编造看不清的内容。"

    try:
        async with httpx.AsyncClient(timeout=60, follow_redirects=True) as client:
            if provider["protocol"] == "openai_chat":
                response = await client.post(
                    _endpoint(endpoint, "/chat/completions"),
                    headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                    json={"model": provider["model"], "messages": [{"role": "user", "content": [{"type": "text", "text": instruction}, {"type": "image_url", "image_url": {"url": f"data:{row['content_type']};base64,{encoded}"}}]}], "max_tokens": 2000},
                )
                if response.status_code >= 400:
                    raise ValueError(f"接口返回 HTTP {response.status_code}")
                return _message_content(response.json()["choices"][0]["message"].get("content"))
            if provider["protocol"] == "anthropic_messages":
                response = await client.post(
                    _endpoint(endpoint, "/messages"),
                    headers={"x-api-key": api_key, "anthropic-version": provider.get("api_version") or "2023-06-01", "content-type": "application/json"},
                    json={"model": provider["model"], "max_tokens": 2000, "messages": [{"role": "user", "content": [{"type": "image", "source": {"type": "base64", "media_type": row["content_type"], "data": encoded}}, {"type": "text", "text": instruction}]}]},
                )
                if response.status_code >= 400:
                    raise ValueError(f"接口返回 HTTP {response.status_code}")
                return "\n".join(str(item.get("text", "")) for item in response.json().get("content", []) if item.get("type") == "text").strip()
            if provider["protocol"] == "gemini_generate":
                response = await client.post(
                    f"{endpoint}/models/{quote(provider['model'], safe='')}:generateContent",
                    params={"key": api_key},
                    json={"contents": [{"parts": [{"text": instruction}, {"inline_data": {"mime_type": row["content_type"], "data": encoded}}]}]},
                )
                if response.status_code >= 400:
                    raise ValueError(f"接口返回 HTTP {response.status_code}")
                parts = response.json()["candidates"][0]["content"]["parts"]
                return "\n".join(str(item.get("text", "")) for item in parts).strip()
    except (httpx.HTTPError, KeyError, IndexError, TypeError, ValueError) as exc:
        raise HTTPException(
            status_code=422,
            detail=f"图片解析失败：{exc}。请确认当前 Provider 的模型支持视觉输入。",
        ) from exc
    raise HTTPException(status_code=422, detail="当前 Provider 协议不支持图片解析")


async def build_attachment_context(attachment_ids: list[str], user_id: int, provider: dict) -> str:
    if not attachment_ids:
        return ""
    unique_ids = list(dict.fromkeys(attachment_ids))
    placeholders = ",".join("?" for _ in unique_ids)
    with db() as conn:
        rows = conn.execute(
            f"SELECT * FROM attachments WHERE user_id = ? AND id IN ({placeholders})",
            (user_id, *unique_ids),
        ).fetchall()
    attachments = {row["id"]: dict(row) for row in rows}
    missing = [attachment_id for attachment_id in unique_ids if attachment_id not in attachments]
    if missing:
        raise HTTPException(status_code=404, detail="部分附件不存在、已删除或无权访问")

    blocks: list[str] = []
    used = 0
    for attachment_id in unique_ids:
        row = attachments[attachment_id]
        content = await _describe_image(row, provider) if row["kind"] == "image" else row["extracted_text"]
        if not content:
            content = "未能从该文件提取可读文本。"
        remaining = MAX_CONTEXT_CHARS - used
        if remaining <= 0:
            break
        content = content[:remaining]
        used += len(content)
        label = "图片视觉解析" if row["kind"] == "image" else "文件内容"
        blocks.append(f"【用户上传{label}：{row['filename']}】\n{content}")
    return "\n\n".join(blocks)
